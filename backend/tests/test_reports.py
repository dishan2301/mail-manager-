import asyncio
import csv
from datetime import datetime, timedelta, timezone
from io import BytesIO, StringIO
from types import SimpleNamespace
from zoneinfo import ZoneInfo

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.api import routes
from app.core.database import Base, get_db
from app.models.entities import AIReviewSetting, Classification, ConversationAIReview, Email, EmailStatus, Employee, Mailbox, SecurityAuditEvent
from app.services import conversation_reviews as reviews


class Clock(datetime):
    @classmethod
    def now(cls, tz=None):
        value = datetime(2026, 11, 1, 1, tzinfo=ZoneInfo("Asia/Kolkata"))
        return value.astimezone(tz) if tz else value.replace(tzinfo=None)


@pytest.fixture
def report_db(monkeypatch):
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    sessions = sessionmaker(bind=engine)
    monkeypatch.setattr(reviews, "SessionLocal", sessions)
    monkeypatch.setattr(reviews, "datetime", Clock)
    monkeypatch.setattr(routes, "datetime", Clock)
    settings = SimpleNamespace(ai_api_key="test-key", app_timezone="Asia/Kolkata", max_report_rows=10000)
    monkeypatch.setattr(reviews, "get_settings", lambda: settings)
    monkeypatch.setattr(routes, "get_settings", lambda: settings)
    monkeypatch.setattr(routes, "mailbox_scope", lambda *_: [1])
    monkeypatch.setattr(routes, "audit", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(reviews, "audit", lambda *_args, **_kwargs: None)
    with sessions() as db:
        db.add_all([Mailbox(id=1, address="one@example.com", display_name="One"), Mailbox(id=2, address="two@example.com", display_name="Two"), Employee(id=1, name="Same Name", email="one@example.com", department="Support"), Employee(id=2, name="Same Name", email="two@example.com", department="Sales"), AIReviewSetting(id=1, enabled=True)])
        db.commit()
    app = FastAPI()
    app.include_router(routes.router)
    def database():
        with sessions() as db:
            yield db
    app.dependency_overrides[get_db] = database
    app.dependency_overrides[routes.viewer.dependency] = lambda: {"user_id": 1, "role": "manager"}
    with TestClient(app) as client:
        yield sessions, client
    engine.dispose()


def mail(db, message_id, when=None, **values):
    values = {"mailbox_id": 1, "message_id": message_id, "conversation_id": message_id, "sender": "customer@example.com", "receiver": "one@example.com", "subject": "Request", "received_time": when or Clock.now(timezone.utc), "folder": "inbox", "classification": Classification.CUSTOMER, "status": EmailStatus.PENDING, "assigned_employee_id": 1, **values}
    row = Email(**values)
    db.add(row)
    db.commit()
    return row


def fake_provider(monkeypatch):
    calls = []
    class Provider:
        async def message_content(self, address, message_id):
            calls.append(message_id)
            return "Private conversation text"
    async def analyze(*args):
        return "Evidence: acknowledgement on the supplied date."
    monkeypatch.setattr(reviews, "_strategy", lambda *_: (Provider(), None))
    monkeypatch.setattr(reviews, "_save_token", lambda *_: None)
    monkeypatch.setattr(reviews, "review_employee_conversations", analyze)
    return calls


def test_work_handling_score_uses_observed_mail_and_missing_speed():
    replied=SimpleNamespace(status=EmailStatus.REPLIED,sla_tier=None,pending_hours=2)
    urgent=SimpleNamespace(status=EmailStatus.PENDING,sla_tier=routes.SLATier.CRITICAL,pending_hours=0)
    result=routes._work_handling_score([replied,urgent],4)
    assert result["score"] is None and result["assessed_points"]==20
    assert [factor["value"] for factor in result["factors"][:2]]==[50,0]
    assert "1h coaching target" in result["factors"][1]["detail"]
    assert routes._work_handling_score([],4)["score"] is None
    assert len(routes._work_handling_score([urgent],None)["factors"])==7
    rubric={key:{"score":80} for key in ("professionalism","detail","clarity","ownership","actionability")}
    assert routes._work_handling_score([replied,urgent],4,[rubric],1)["score"]==69


def test_calendar_boundaries_and_weekly_review_at_month_change(report_db, monkeypatch):
    sessions, client = report_db
    assert routes._report_cutoff("daily") == datetime(2026, 10, 31, 18, 30, tzinfo=timezone.utc)
    assert routes._report_cutoff("monthly") == routes._report_cutoff("daily")
    with sessions() as db:
        mail(db, "previous-month", Clock.now(timezone.utc) - timedelta(days=2))
        mail(db, "current-month")
    calls = fake_provider(monkeypatch)
    assert asyncio.run(reviews.review_new_or_changed_conversations())["reviewed"] == 2
    assert len(calls) == 2
    assert client.get("/api/reports/weekly/employees/1/ai-review").json()["reviewed_count"] == 2
    assert client.get("/api/reports/monthly/employees/1/ai-review").json()["reviewed_count"] == 1
    result = client.get("/api/reports/monthly?dimension=department").json()
    assert result["period_start"] == "2026-10-31T18:30:00+00:00"
    assert result["latest_received_at"] == Clock.now(timezone.utc).isoformat()


def test_cache_skips_unchanged_and_detects_outgoing_reply(report_db, monkeypatch):
    sessions, client = report_db
    with sessions() as db:
        mail(db, "thread")
    calls = fake_provider(monkeypatch)
    assert asyncio.run(reviews.review_new_or_changed_conversations())["reviewed"] == 1
    assert asyncio.run(reviews.review_new_or_changed_conversations())["reviewed"] == 0
    assert calls == ["thread"]
    with sessions() as db:
        mail(db, "outgoing", Clock.now(timezone.utc) + timedelta(minutes=5), conversation_id="thread", folder="sentitems", sender="one@example.com")
    result = client.get("/api/reports/monthly/employees/1/ai-review").json()
    assert result["pending_count"] == 1 and result["reviews"][0]["analysis"] is None
    assert calls == ["thread"]  # Viewing a report never fetches mail or calls AI.
    assert asyncio.run(reviews.review_new_or_changed_conversations())["reviewed"] == 1
    assert client.get("/api/reports/monthly/employees/1/ai-review").json()["reviewed_count"] == 1


def test_consent_and_failure_retry(report_db, monkeypatch):
    sessions, client = report_db
    calls = fake_provider(monkeypatch)
    with sessions() as db:
        mail(db, "thread")
        db.get(AIReviewSetting, 1).enabled = False
        db.commit()
    assert "skipped" in asyncio.run(reviews.review_new_or_changed_conversations())
    assert not calls
    with sessions() as db:
        db.get(AIReviewSetting, 1).enabled = True
        db.commit()
    async def fail(*_):
        raise reviews.WritingAssistantError("secret provider error")
    monkeypatch.setattr(reviews, "review_employee_conversations", fail)
    assert asyncio.run(reviews.review_new_or_changed_conversations())["failed"] == 1
    result = client.get("/api/reports/monthly/employees/1/ai-review").json()
    assert result["failed_count"] == 1 and "secret" not in str(result)
    fake_provider(monkeypatch)
    assert asyncio.run(reviews.review_new_or_changed_conversations())["reviewed"] == 0
    with sessions() as db:
        db.scalar(select(ConversationAIReview)).analyzed_at = Clock.now(timezone.utc)-timedelta(minutes=5)
        db.commit()
    assert asyncio.run(reviews.review_new_or_changed_conversations())["reviewed"] == 1


def test_totals_employee_identity_scoping_and_exports(report_db, monkeypatch):
    sessions, client = report_db
    with sessions() as db:
        mail(db, "pending")
        mail(db, "replied", status=EmailStatus.REPLIED, pending_hours=2)
        mail(db, "ignored", status=EmailStatus.IGNORED)
        mail(db, "other-employee", assigned_employee_id=2)
        mail(db, "private-mailbox", mailbox_id=2)
    endpoint = "/api/reports/monthly"
    rows = client.get(endpoint + "?dimension=employee").json()["rows"]
    assert len(rows) == 2 and {row["employee_id"] for row in rows} == {1, 2}
    row = next(row for row in rows if row["employee_id"] == 1)
    assert (row["total_received"], row["replied"], row["not_replied"], row["ignored"]) == (3, 1, 1, 1)
    assert row["reply_rate"] == 50 and row["average_reply_hours"] == 2
    assert {row["dimension"] for row in client.get(endpoint + "?dimension=department").json()["rows"]} == {"Support", "Sales"}
    assert client.get(endpoint + "?dimension=customer").json()["rows"][0]["total_received"] == 4
    assert client.get(endpoint + "?dimension=customer&employee_id=2").json()["rows"][0]["total_received"] == 1
    response = client.get(endpoint + "/export?dimension=employee&format=csv")
    exported = list(csv.DictReader(StringIO(response.content.decode("utf-8-sig"))))
    assert len(exported) == 2 and sum(int(row["total_received"]) for row in exported) == 4
    response = client.get(endpoint + "/export?dimension=employee&format=csv&department=Support&employee_id=1")
    assert len(list(csv.DictReader(StringIO(response.content.decode("utf-8-sig"))))) == 1
    response = client.get(endpoint + "/export?dimension=employee&format=csv&department=Sales&employee_id=1")
    assert len(list(csv.DictReader(StringIO(response.content.decode("utf-8-sig"))))) == 0
    response = client.get(endpoint + "/export?dimension=employee&format=xlsx")
    workbook = load_workbook(BytesIO(response.content))
    assert workbook.active.max_row == 3
    assert client.get(endpoint + "/export?dimension=employee&format=pdf").content.startswith(b"%PDF")
    assert client.patch("/api/reports/ai-review/status", json={"enabled": True}).status_code in {401, 403}
    monkeypatch.setattr(routes, "mailbox_scope", lambda *_: [])
    no_access = client.get(endpoint + "?dimension=employee").json()
    assert no_access["rows"] == [] and no_access["latest_received_at"] is None
    assert client.get(endpoint + "/employees/1/ai-review").json()["conversation_count"] == 0


def test_daily_run_covers_all_changed_conversations_and_paginates(report_db, monkeypatch):
    sessions, client = report_db
    with sessions() as db:
        for index in range(101):
            mail(db, f"thread-{index}")
    fake_provider(monkeypatch)
    assert asyncio.run(reviews.review_new_or_changed_conversations())["reviewed"] == 101
    result = client.get("/api/reports/monthly/employees/1/ai-review").json()
    assert (result["reviewed_count"], result["pending_count"], result["total_pages"]) == (101, 0, 2)
    assert asyncio.run(reviews.review_new_or_changed_conversations())["reviewed"] == 0
    with sessions() as db:
        assert len(list(db.scalars(select(ConversationAIReview)))) == 101


def test_shared_thread_has_separate_employee_findings(report_db, monkeypatch):
    sessions, client = report_db
    with sessions() as db:
        mail(db, "first", conversation_id="shared", assigned_employee_id=1)
        mail(db, "second", conversation_id="shared", assigned_employee_id=2)
    fake_provider(monkeypatch)
    assert asyncio.run(reviews.review_new_or_changed_conversations())["reviewed"] == 2
    for employee_id in (1, 2):
        assert client.get(f"/api/reports/monthly/employees/{employee_id}/ai-review").json()["reviewed_count"] == 1


def test_unsynced_reply_waits_and_assistant_events_invalidate_cache(report_db, monkeypatch):
    sessions, client = report_db
    calls = fake_provider(monkeypatch)
    with sessions() as db:
        incoming = mail(db, "thread")
        email_id = incoming.id
    assert asyncio.run(reviews.review_new_or_changed_conversations())["reviewed"] == 1
    with sessions() as db:
        db.get(Email, email_id).replied_at = Clock.now(timezone.utc) + timedelta(minutes=5)
        db.commit()
    assert client.get("/api/reports/monthly/employees/1/ai-review").json()["pending_count"] == 1
    assert asyncio.run(reviews.review_new_or_changed_conversations())["reviewed"] == 0
    assert calls == ["thread"]
    with sessions() as db:
        mail(db, "sent", Clock.now(timezone.utc) + timedelta(minutes=5), folder="sentitems", conversation_id="thread", sender="one@example.com")
    assert asyncio.run(reviews.review_new_or_changed_conversations())["reviewed"] == 1
    with sessions() as db:
        db.add(SecurityAuditEvent(action="email.reply.assist", outcome="success", object_type="email", object_id=str(email_id), created_at=Clock.now(timezone.utc) + timedelta(minutes=10), details={"mode":"rewrite", "has_employee_draft":True, "input_word_count":25}))
        db.commit()
        evidence = reviews._ai_assistance(db, [db.get(Email, email_id)])
        assert evidence[0]["has_employee_draft"] is True
        assert evidence[0]["input_word_count"] == 25 and evidence[0]["sent_usage_verified"] is False
    assert client.get("/api/reports/monthly/employees/1/ai-review").json()["pending_count"] == 1
    assert asyncio.run(reviews.review_new_or_changed_conversations())["reviewed"] == 1


def test_ai_review_migrations_preserve_existing_findings(report_db):
    import importlib
    from alembic.migration import MigrationContext
    from alembic.operations import Operations
    from sqlalchemy import inspect
    sessions, _ = report_db
    engine = sessions.kw["bind"]
    old = importlib.import_module("migrations.versions.0009_daily_ai_reviews")
    new = importlib.import_module("migrations.versions.0010_employee_conversation_reviews")
    with engine.begin() as connection:
        context = MigrationContext.configure(connection)
        with Operations.context(context):
            old.upgrade()  # Fresh databases already contain current metadata.
            new.upgrade()
            new.downgrade()  # Reproduce the previous one-review-per-thread schema.
        connection.execute(ConversationAIReview.__table__.insert().values(mailbox_id=1, conversation_key="thread", employee_id=1, received_at=Clock.now(timezone.utc), last_message_at=Clock.now(timezone.utc), status="ready", analysis="Saved evidence"))
        with Operations.context(context):new.upgrade()
        assert connection.scalar(select(ConversationAIReview.analysis)) == "Saved evidence"
        columns = inspect(connection).get_unique_constraints("conversation_ai_reviews")[0]["column_names"]
        assert columns == ["mailbox_id", "conversation_key", "employee_id"]


@pytest.mark.parametrize("mode,body,has_draft", [("draft", "", False), ("rewrite", "My own answer", True)])
def test_assistant_request_records_draft_presence_without_storing_text(report_db, monkeypatch, mode, body, has_draft):
    sessions, client = report_db
    with sessions() as db:
        email_id = mail(db, "request").id
    client.app.dependency_overrides[routes.admin.dependency] = lambda: {"user_id": 1, "role": "admin"}
    class Provider:
        async def message_content(self, *_):return "Customer context"
    async def assist(*_):return "Generated suggestion"
    events = []
    monkeypatch.setattr(routes, "_email_strategy", lambda *_: (Provider(), None))
    monkeypatch.setattr(routes, "_save_refreshed_mailbox_token", lambda *_: None)
    monkeypatch.setattr(routes, "suggest_reply", assist)
    monkeypatch.setattr(routes, "audit", lambda *args, **kwargs: events.append(args))
    response = client.post(f"/api/emails/{email_id}/rewrite", json={"mode":mode, "body":body})
    assert response.status_code == 200
    evidence = events[-1][-1]
    assert evidence["has_employee_draft"] is has_draft and evidence["input_word_count"] == len(body.split())
    assert "body" not in evidence


def test_pdf_keeps_all_fields_and_csv_neutralizes_formula_senders(report_db, tmp_path):
    import subprocess
    sessions, client = report_db
    with sessions() as db:
        mail(db, "malicious-sender", sender="=HYPERLINK(\"example\")", status=EmailStatus.REPLIED, pending_hours=2)
    response = client.get("/api/reports/monthly/export?dimension=customer&format=csv")
    row = next(csv.DictReader(StringIO(response.content.decode("utf-8-sig"))))
    assert row["dimension"].startswith("'=HYPERLINK")
    pdf = tmp_path / "report.pdf"
    pdf.write_bytes(client.get("/api/reports/monthly/export?dimension=customer&format=pdf").content)
    content = subprocess.run(["pdftotext", str(pdf), "-"], capture_output=True, text=True, check=True).stdout
    assert "Average Reply Hours:" in content and "Reply Rate:" in content and "Total Received:" in content


def test_employees_without_period_mail_remain_visible(report_db, monkeypatch):
    sessions, client = report_db
    with sessions() as db:
        mail(db, "older", Clock.now(timezone.utc) - timedelta(days=40))
    rows = client.get("/api/reports/monthly?dimension=employee").json()["rows"]
    assert len(rows) == 1 and rows[0]["employee_id"] == 1 and rows[0]["total_received"] == 0
    monkeypatch.setattr(routes, "mailbox_scope", lambda *_: None)
    rows = client.get("/api/reports/monthly?dimension=employee").json()["rows"]
    assert {row["employee_id"] for row in rows} == {1, 2}
    assert all(row["total_received"] == 0 for row in rows)


def test_long_conversation_reviews_every_section_before_ready(report_db, monkeypatch):
    import copy
    sessions, client = report_db
    with sessions() as db:
        mail(db, "large")
    seen = []
    class Provider:
        async def message_content(self, *_):return "x" * 81001
    async def analyze(*args):
        seen.append(copy.deepcopy(args[-1][0]))
        return f"Integrated findings through section {len(seen)}"
    monkeypatch.setattr(reviews, "_strategy", lambda *_: (Provider(), None))
    monkeypatch.setattr(reviews, "_save_token", lambda *_: None)
    monkeypatch.setattr(reviews, "review_employee_conversations", analyze)
    assert asyncio.run(reviews.review_new_or_changed_conversations())["reviewed"] == 1
    assert len(seen) == 3
    assert sum(len(message["body"]) for section in seen for message in section["messages"]) == 81001
    assert seen[0]["prior_review"] is None
    assert seen[2]["prior_review"] == "Integrated findings through section 2"
    result = client.get("/api/reports/monthly/employees/1/ai-review").json()
    assert result["reviews"][0]["analysis"] == "Integrated findings through section 3"


def test_admin_consent_is_explicit_and_requires_configured_provider(report_db):
    sessions, client = report_db
    client.app.dependency_overrides[routes.admin.dependency] = lambda: {"user_id": 1, "role": "admin"}
    assert client.patch("/api/reports/ai-review/status", json={"enabled":"true"}).status_code == 422
    response = client.patch("/api/reports/ai-review/status", json={"enabled":True})
    assert response.status_code == 200 and response.json()["enabled"] is True
    with sessions() as db:
        setting = db.get(AIReviewSetting, 1)
        assert setting.approved_by == 1 and setting.approved_at is not None
    assert client.patch("/api/reports/ai-review/status", json={"enabled":False}).json()["enabled"] is False
    settings = routes.get_settings()
    settings.ai_api_key = ""
    assert client.patch("/api/reports/ai-review/status", json={"enabled":True}).status_code == 503
    assert client.get("/api/reports/ai-review/status").json()["enabled"] is False


def test_continuous_review_schedule_starts_immediately_and_checks_every_30_seconds(monkeypatch):
    from app.services import scheduler
    monkeypatch.setattr(scheduler, "get_settings", lambda: SimpleNamespace(app_timezone="Asia/Kolkata", summary_hour=18, sync_interval_seconds=10))
    async def operation():pass
    jobs = scheduler.build_scheduler(operation, operation, operation)
    before = datetime.now(timezone.utc)
    job = jobs.get_job("conversation-ai-review")
    assert job.trigger.interval.total_seconds() == 30
    assert abs((job.next_run_time-before).total_seconds()) < 2
    assert job.max_instances == 1 and job.coalesce is True


def test_five_conversations_are_reviewed_independently(report_db, monkeypatch):
    sessions, _ = report_db
    with sessions() as db:
        for index in range(5):
            mail(db, f"thread-{index}", subject=f"Request {index}", assigned_employee_id=1 if index<3 else 2)
    seen = []
    fake_provider(monkeypatch)
    async def analyze(settings, employee, department, period, conversations):
        assert len(conversations) == 1
        assert conversations[0]["prior_review"] is None
        seen.append((employee,department,conversations[0]["subject"]))
        return f"Evidence for {conversations[0]['subject']}"
    monkeypatch.setattr(reviews, "review_employee_conversations", analyze)
    assert asyncio.run(reviews.review_new_or_changed_conversations())["reviewed"] == 5
    assert len(seen) == 5
    assert {subject for _,_,subject in seen} == {f"Request {index}" for index in range(5)}
    assert sum(department=="Support" for _,department,_ in seen) == 3
    assert sum(department=="Sales" for _,department,_ in seen) == 2


def test_reply_arriving_during_review_is_queued_for_an_update(report_db, monkeypatch):
    sessions, client = report_db
    with sessions() as db:mail(db, "thread")
    fake_provider(monkeypatch)
    async def analyze(*_):
        with sessions() as db:
            mail(db, "new-reply", Clock.now(timezone.utc)+timedelta(minutes=1), conversation_id="thread", folder="sentitems", sender="one@example.com")
        return "Findings before the new reply arrived"
    monkeypatch.setattr(reviews, "review_employee_conversations", analyze)
    assert asyncio.run(reviews.review_new_or_changed_conversations())["reviewed"] == 1
    result=client.get("/api/reports/monthly/employees/1/ai-review").json()
    assert result["pending_count"] == 1 and result["reviews"][0]["analysis"] is None
    fake_provider(monkeypatch)
    assert asyncio.run(reviews.review_new_or_changed_conversations())["reviewed"] == 1
    assert client.get("/api/reports/monthly/employees/1/ai-review").json()["reviewed_count"] == 1
